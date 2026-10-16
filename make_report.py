"""Build the Experiment 9 lab report (.docx) in the course template, then PDF via LibreOffice.

The Experiment 7 report is used as the template so the header lines, metadata
table, fonts and margins match the required format exactly.

    ~/Downloads/.venv_gpt2/bin/python make_report.py
"""
from __future__ import annotations

import copy
import subprocess
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parent
SHOTS = BASE / "screenshots"
DOWNLOADS = BASE.parent
TEMPLATE = DOWNLOADS / "Anson Dias_2307040_Experiment7_vLLM_Serving.docx"
OUT = DOWNLOADS / "Anson Dias_2307040_Experiment9_AI_Email_Agent.docx"

META = {
    "Experiment No.": "9",
    "Title": "Building an AI Agent: an Autonomous Email Assistant "
             "(LLM + Memory + Tools + Orchestration)",
    "Lab Outcome Achieved": (
        "Designed and built an autonomous AI email agent that combines a local LLM "
        "(Qwen2.5-7B via Ollama) for intent understanding and planning, tool calling for "
        "contact lookup and Gmail SMTP delivery, short- and long-term memory for follow-up "
        "instructions, and code-enforced guardrails (address verification and human "
        "confirmation) around the irreversible send action."),
    "Date of Performance": "16/10/2026",
    "Date of submission": "16/10/2026",
    "APP ID": "2307040",
    "Name of the Student": "Anson Dias",
}

doc = docx.Document(str(TEMPLATE))
body = doc.element.body

# keep: Class line, Subject line, metadata table, blank paragraph, section properties
for el in list(body)[4:]:
    if not el.tag.endswith("sectPr"):
        body.remove(el)
table = doc.tables[0]
for row in table.rows:
    label = row.cells[0].text.strip()
    if label in META:
        cell = row.cells[1]
        for p in cell.paragraphs[1:]:
            p._element.getparent().remove(p._element)
        para = cell.paragraphs[0]
        for r in para.runs[1:]:
            r._element.getparent().remove(r._element)
        if para.runs:
            para.runs[0].text = META[label]
        else:
            para.add_run(META[label])


# ---------------------------------------------------------------- helpers
def _space(p, after):
    p.paragraph_format.space_after = Pt(after)
    return p


def section(title):
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(13)
    return _space(p, 6)


def heading(text):
    p = doc.add_paragraph()
    p.add_run(text).bold = True
    p.paragraph_format.keep_with_next = True
    return _space(p, 3)


def para(text, after=6):
    """Paragraph with **bold** spans."""
    p = doc.add_paragraph()
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part:
            p.add_run(part).bold = (i % 2 == 1)
    return _space(p, after)


def bullet(text):
    p = para(text, after=3)
    p.style = doc.styles["List Bullet"]
    return p


def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Courier New"
    r.font.size = Pt(9.5)
    p.paragraph_format.left_indent = Inches(0.25)
    return _space(p, 6)


def figure(name, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(SHOTS / name), width=Inches(width))
    _space(p, 2)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.italic = True
    r.font.size = Pt(9.5)
    _space(c, 10)


def grid(rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = table.style
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.width = Inches(widths[j])
            cell.paragraphs[0].text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            run.bold = (i == 0)
    doc.add_paragraph()
    return t


# ---------------------------------------------------------------- THEORY
section("Theory:")
para("A chatbot only produces text; an **AI agent** pursues a goal by deciding what to do, "
     "doing it through tools, observing the result and continuing until the goal is met. "
     "A useful working definition is **AI Agent = LLM (reasoning/planning) + Memory "
     "(context/state) + Tools (actions) + Orchestration**. The aim of this experiment is "
     "to build such an agent end-to-end: an email assistant that accepts an instruction "
     "such as “Send an email to Dr. Meher saying my Experiment 9 agent is complete” "
     "and autonomously understands it, looks up the recipient, drafts the email, asks for "
     "confirmation, sends it through Gmail and remembers the interaction for follow-ups.")

heading("1. The four components of an agent")
grid([
    ["Component", "Role", "This implementation"],
    ["LLM", "Understands the request, plans the next step, chooses tools, writes text",
     "Qwen2.5-7B-Instruct (Q4_K_M) served locally by Ollama, temperature 0 (llm.py)"],
    ["Memory", "Keeps context within a conversation and state across sessions",
     "Short-term: last 12 turns replayed to the LLM; long-term: every email with its status "
     "in data/memory.json (memory.py)"],
    ["Tools", "Let the agent act on the world and fetch facts it does not know",
     "lookup_contact, save_contact, recall_recent_emails, send_email, resend_email (tools.py)"],
    ["Orchestration", "Runs the loop, passes results back, enforces limits and safety",
     "understand → ReAct tool loop → reply, step limit, stall recovery, guardrails (agent.py)"],
], [1.1, 2.4, 2.8])

heading("2. Tool calling and the ReAct loop")
para("Modern LLMs support **function (tool) calling**: each tool is described to the model "
     "by a name, a description and a JSON schema of its arguments. Instead of answering in "
     "prose the model can emit a structured call such as lookup_contact(name=\"Dr. Meher\"). "
     "The orchestrator executes the real function and appends its result (the "
     "observation) to the conversation, and the model reasons again with the new "
     "information. This Reason → Act → Observe cycle is the **ReAct** pattern. "
     "The loop ends when the model replies in plain text, and a hard step limit (8) prevents "
     "runaway loops:")
code("messages = [system_prompt, *recent_turns, current_task_note]\n"
     "for step in range(MAX_AGENT_STEPS):\n"
     "    reply = llm.chat(messages, tools=TOOL_SCHEMAS)\n"
     "    if not reply.tool_calls:  return reply.content        # final answer\n"
     "    for call in reply.tool_calls:\n"
     "        result = toolbox.run(call.name, call.arguments)   # guardrails live here\n"
     "        messages.append({\"role\": \"tool\", \"content\": result})")

heading("3. Intent and entity extraction with structured output")
para("Step 1 of the agent turns free text into structured data. Ollama supports "
     "**grammar-constrained decoding**: when a JSON schema is passed in the format field, "
     "token sampling is restricted so the output is always valid JSON matching the schema. "
     "The schema used here extracts intent (send_email, resend_previous, provide_info, "
     "query_history, cancel, other), recipients exactly as named, a subject, key points, the "
     "tone (formal/friendly/neutral/urgent) and whether the message refers to an earlier "
     "email. The parsed request is pinned into the next stage as a “current task” "
     "note so the planner acts on this instruction and not on older ones in the history.")

heading("4. Memory: short-term and long-term")
para("**Short-term memory** is the conversation buffer: the most recent turns are replayed "
     "in the model’s context window on every request, so a reply like “It’s "
     "john.doe@example.com” is understood as the answer to the agent’s question. "
     "**Long-term memory** is a persistent store outside the model: every drafted, sent, "
     "declined or failed email is written to data/memory.json and read back with the "
     "recall_recent_emails tool. A follow-up such as “Send the same email to Sarah "
     "too” needs both: the reference “the same email” is resolved to a stored "
     "email, and resend_email reuses its subject and body word-for-word (only the greeting "
     "changes), so the copy cannot drift from what was originally approved. Because the "
     "store is on disk, follow-ups work even after the program is restarted.")

heading("5. Guardrails for an irreversible action")
para("LLMs can **hallucinate**: asked to email “John”, a model may invent a "
     "plausible address such as john@gmail.com. Sending an email cannot be undone, so the "
     "safety rules are enforced in code at the tool boundary rather than trusted to the "
     "prompt:")
bullet("**Address verification (grounding):** send_email, resend_email and save_contact "
       "accept an address only if it is in the contacts store or was literally typed by the "
       "user in the conversation. Anything else returns BLOCKED to the model, which must "
       "look the person up or ask.")
bullet("**Human-in-the-loop confirmation:** the full draft (From, To, Subject, body) is "
       "shown and nothing is sent without an explicit “y”. “n” cancels; "
       "any other text is returned to the model as revision feedback. Setting "
       "CONFIRM_BEFORE_SEND=false switches to direct sending, but verification still applies.")
bullet("**Idempotency:** the same email cannot be sent twice for one instruction, and a "
       "declined recipient is not retried.")
para("Each guardrail returns a clear observation (BLOCKED, NOT_SENT, ALREADY_SENT) that the "
     "model can reason about, so a blocked action leads to recovery (asking the user) "
     "rather than a crash.")

heading("6. Sending email: SMTP and Gmail App Passwords")
para("Email is delivered with the **Simple Mail Transfer Protocol (SMTP)**. Python’s "
     "smtplib opens an implicitly encrypted TLS connection to smtp.gmail.com on port 465, "
     "authenticates and submits an RFC 5322 message built with email.message.EmailMessage. "
     "Google does not accept the normal account password from programs. Instead, after "
     "2-Step Verification is enabled, the account owner generates a **16-character App "
     "Password** for one application. It can be revoked at any time without changing the "
     "main password. Here it is entered once through a hidden prompt "
     "(setup_credentials.py), checked against Gmail and stored in a .env file with "
     "owner-only permissions. The LLM never sees the credential: it can only ask the "
     "send tool to send, and the tool reads the secret from configuration.")

heading("7. System design and execution environment")
para("Figure 0 shows the complete agent. Everything except email delivery runs locally "
     "on an Apple M1 MacBook (8 GB RAM): the LLM is qwen2.5:7b (4.7 GB, 4-bit) on Ollama, "
     "and the agent uses only the Python 3.14 standard library (urllib, smtplib, imaplib, "
     "json). No cloud LLM or API key is used.")
figure("00_architecture.png", "Fig. 0 — Architecture of the AI email agent: LLM, "
       "orchestration loop, tools, guardrails, memory and delivery.")

# ---------------------------------------------------------------- OUTPUT
section("Output:")
exec((BASE / "report_output.py").read_text())  # Output section text (written after the live run)

# ---------------------------------------------------------------- save + PDF
doc.save(str(OUT))
print("saved", OUT)
subprocess.run(["/opt/homebrew/bin/soffice", "--headless", "--convert-to", "pdf",
                "--outdir", str(DOWNLOADS), str(OUT)], check=True,
               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print("saved", OUT.with_suffix(".pdf"))
